#!/usr/bin/env python3
"""
Reproduce the WT cAMP Emax normalization from the raw dose-response data in
transcriptional_drc.xlsx and plot the normalized curves as a single overlay.

Pipeline (reproduces WT_cAMP_vs_TRUPATH_All_Drugs Emax exactly):
  1. Parse the Prism multi-block sheet -> tidy (ligand, logdose, response, flagged).
  2. Fit each ligand independently: 3-param log-logistic, Hill slope FIXED = 1,
     asterisked/flagged points excluded.
        response = Bottom + (Top - Bottom) / (1 + 10^(LogEC50 - logdose))
     Top = high-dose plateau (max inhibition), Bottom = baseline (~1.0).
  3. Span = Top - Bottom ; Emax(% DAMGO) = Span_ligand / Span_DAMGO * 100.
  4. Normalize to "% DAMGO effect": effect% = (Bottom_i - response) / |Span_DAMGO| * 100
     -> each curve rises from 0% (baseline) to its Emax plateau; DAMGO tops at 100%.
"""
import matplotlib
matplotlib.use("Agg")
import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import numpy as np
import openpyxl
from scipy.optimize import curve_fit

XLSX = "transcriptional_drc.xlsx"
OUT  = "camp_drc_curves.pdf"

# ── house style ──────────────────────────────────────────────────────────────
mpl.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Helvetica", "Arial", "DejaVu Sans"],
    "pdf.fonttype": 42, "ps.fonttype": 42,
    "font.size": 6, "axes.labelsize": 6, "axes.titlesize": 6,
    "xtick.labelsize": 6, "ytick.labelsize": 6, "legend.fontsize": 6,
    "text.color": "black", "axes.labelcolor": "black",
    "xtick.color": "black", "ytick.color": "black", "axes.edgecolor": "black",
    "axes.linewidth": 0.5, "xtick.major.width": 0.5, "ytick.major.width": 0.5,
    "xtick.major.size": 2, "ytick.major.size": 2,
    "xtick.direction": "out", "ytick.direction": "out",
})

ligand_palette = {
    "DAMGO":"#1F77B4","Fentanyl":"#FF7F0E","Morphine":"#2CA02C","Naloxone":"#D62728",
    "Buprenorphine":"#9467BD","Methadone":"#8C564B","PZM21":"#E377C2","Carfentanil":"#17BECF",
    "Nalbuphine":"#BCBD22","Butorphanol":"#7F7F7F","TRV130":"#AEC7E8","MP":"#FFBB78",
}

# ── parse the Prism multi-block layout ───────────────────────────────────────
BLOCKS = {2:"DAMGO",11:"Fentanyl",20:"Methadone",29:"Buprenorphine",38:"Morphine",
          47:"Naloxone",56:"PZM21",65:"Uninduced cells",74:"TRV130",83:"MP",
          92:"Carfentanil",101:"Nalbuphine",110:"Butorphanol"}
def parse():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["Sheet1"]
    starts = sorted(BLOCKS)
    data = {}  # ligand -> list of (logdose, response, flagged)
    for i, c in enumerate(starts):
        nxt = starts[i+1] if i+1 < len(starts) else 119
        name = BLOCKS[c]
        pts = []
        for r in range(2, ws.max_row+1):
            xa = ws.cell(r, 1).value
            if xa in (None, ""):
                continue
            for cc in range(c, nxt):
                v = ws.cell(r, cc).value
                if v in (None, ""):
                    continue
                flagged = isinstance(v, str) and "*" in v
                try:
                    fv = float(str(v).replace("*", "").strip())
                except ValueError:
                    continue
                pts.append((float(xa), fv, flagged))
        data[name] = pts
    return data

# ── fit: 3-param, Hill = 1 ───────────────────────────────────────────────────
def model(x, top, bot, le):
    return bot + (top - bot) / (1 + 10.0**(le - x))

def fit_ligand(pts):
    good = [(x, y) for (x, y, fl) in pts if not fl]
    x = np.array([p[0] for p in good]); y = np.array([p[1] for p in good])
    p, _ = curve_fit(model, x, y, p0=[np.min(y), np.max(y), -8.0], maxfev=40000)
    return dict(top=p[0], bot=p[1], le=p[2])

def main():
    data = parse()
    fits = {lig: fit_ligand(data[lig]) for lig in ligand_palette}
    damgo_span = fits["DAMGO"]["bot"] - fits["DAMGO"]["top"]   # positive, ~0.616

    def to_effect(resp, bot):
        return (bot - resp) / damgo_span * 100.0

    # Emax table + legend order (descending Emax)
    emax = {lig: (fits[lig]["bot"] - fits[lig]["top"]) / damgo_span * 100.0
            for lig in ligand_palette}
    order = ["DAMGO"] + [l for l in sorted(ligand_palette, key=lambda l: -emax[l])
                         if l != "DAMGO"]
    print("Emax (% DAMGO), reproduced:")
    for lig in order:
        print(f"  {lig:14s} {emax[lig]:6.1f}   (LogEC50 {fits[lig]['le']:7.2f})")

    # ── plot ────────────────────────────────────────────────────────────────
    fig = plt.figure()
    fig.set_size_inches(78/25.4, 46/25.4)
    ax = fig.add_axes([0.115, 0.165, 0.55, 0.80])

    bar_hi, bar_lo = [], []
    for lig in order:
        f = fits[lig]
        is_damgo = lig == "DAMGO"
        col = "black" if is_damgo else ligand_palette[lig]   # DAMGO drawn black
        z = 10 if is_damgo else 2                            # DAMGO on top layer
        # per-dose mean + SEM (flagged excluded), then normalize
        good = [(x, y) for (x, y, fl) in data[lig] if not fl]
        doses = sorted(set(x for x, _ in good))
        mx, my, se = [], [], []
        for d in doses:
            ys = np.array([y for x, y in good if x == d])
            m = ys.mean()
            s = ys.std(ddof=1) / np.sqrt(len(ys)) if len(ys) > 1 else 0.0
            mx.append(d); my.append(to_effect(m, f["bot"])); se.append(s / damgo_span * 100)
        # fitted curve, clipped to 1/2 log beyond the ligand's first/last data point
        xlo = max(-14, min(doses) - 0.5); xhi = min(-4.5, max(doses) + 0.5)
        xx = np.linspace(xlo, xhi, 300)
        yy = to_effect(model(xx, f["top"], f["bot"], f["le"]), f["bot"])
        ax.plot(xx, yy, color=col, lw=0.9, zorder=z)
        ax.errorbar(mx, my, yerr=se, fmt="o", ms=2.2, mfc=col, mec="none",
                    ecolor=col, elinewidth=0.5, capsize=1, capthick=0.5,
                    zorder=z + 1)
        bar_hi += [m + s for m, s in zip(my, se)]
        bar_lo += [m - s for m, s in zip(my, se)]

    ax.axhline(0, color="grey", lw=0.4, ls=(0, (3, 3)), zorder=0)
    ax.set_xlim(-14, -4.5)
    ax.set_ylim(min(-10, min(bar_lo) - 3), max(bar_hi) + 4)
    ax.set_xticks(range(-14, -4, 1))
    ax.set_xticklabels([str(v) if v % 2 == 0 else "" for v in range(-14, -4, 1)])
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.set_xlabel("Log[agonist] (M)")
    ax.set_ylabel("Activity (% DAMGO)")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)

    # dot-only legend handles (no error-bar lines), DAMGO first
    handles = [Line2D([0], [0], marker="o", linestyle="none", markersize=2.6,
                      markerfacecolor=("black" if l == "DAMGO" else ligand_palette[l]),
                      markeredgecolor="none", label=l) for l in order]
    ax.legend(handles=handles, loc="center left", bbox_to_anchor=(1.02, 0.5),
              handlelength=1.0, handletextpad=0.4, labelspacing=0.35,
              borderpad=0.0, frameon=False)

    fig.savefig(OUT, transparent=True)
    print("Saved:", OUT)

if __name__ == "__main__":
    main()
