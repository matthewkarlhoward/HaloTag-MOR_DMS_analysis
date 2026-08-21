#!/usr/bin/env python3
"""
Composite: transcriptional (cAMP) DRC left, TruPath Gi1 DRC right, one shared
legend covering all 14 ligands. Each panel keeps its own validated normalization:
  * cAMP  -> 3-param Hill=1 fit, effect% = (Bottom_i - resp)/DAMGO_span*100
            (see camp_drc_curves.py)
  * TruPath-> activity = 1 - norm; % of DAMGO's PEAK mean activity
            (see trupath_drc_curves.py)
DAMGO is drawn black and on top in both panels. cAMP has 12 ligands; TruPath 14
(adds Naltrexone, C6Guano) -- the legend lists all 14.
"""
import matplotlib
matplotlib.use("Agg")
import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import re
import numpy as np
import openpyxl
from scipy.optimize import curve_fit

CAMP_XLSX = "transcriptional_drc.xlsx"
TP_PZFX   = "20260117_mor_wt_trupath_compilation.pzfx"   # raw TruPath data (first tab)
OUT       = "camp_trupath_composite.pdf"

mpl.rcParams.update({
    "font.family": "sans-serif", "font.sans-serif": ["Helvetica", "Arial", "DejaVu Sans"],
    "pdf.fonttype": 42, "ps.fonttype": 42,
    "font.size": 6, "axes.labelsize": 6, "axes.titlesize": 6,
    "xtick.labelsize": 6, "ytick.labelsize": 6, "legend.fontsize": 6,
    "text.color": "black", "axes.labelcolor": "black",
    "xtick.color": "black", "ytick.color": "black", "axes.edgecolor": "black",
    "axes.linewidth": 0.5, "xtick.major.width": 0.5, "ytick.major.width": 0.5,
    "xtick.major.size": 2, "ytick.major.size": 2, "xtick.direction": "out", "ytick.direction": "out",
})

FULL_PALETTE = {
    "DAMGO":"#1F77B4","Fentanyl":"#FF7F0E","Morphine":"#2CA02C","Naloxone":"#D62728",
    "Buprenorphine":"#9467BD","Methadone":"#8C564B","PZM21":"#E377C2","Carfentanil":"#17BECF",
    "Nalbuphine":"#BCBD22","Butorphanol":"#7F7F7F","TRV130":"#AEC7E8","MP":"#FFBB78",
    "Naltrexone":"#843C39","C6Guano":"#5254A3",
}
def color_for(l): return "black" if l == "DAMGO" else FULL_PALETTE[l]
# legend order: DAMGO first, then by TruPath Emax (has all 14)
LEGEND_ORDER = ["DAMGO","Morphine","Carfentanil","Fentanyl","C6Guano","Methadone","PZM21",
                "TRV130","MP","Butorphanol","Buprenorphine","Nalbuphine","Naltrexone","Naloxone"]

# ── cAMP panel ───────────────────────────────────────────────────────────────
CAMP_BLOCKS = {2:"DAMGO",11:"Fentanyl",20:"Methadone",29:"Buprenorphine",38:"Morphine",
               47:"Naloxone",56:"PZM21",65:"Uninduced cells",74:"TRV130",83:"MP",
               92:"Carfentanil",101:"Nalbuphine",110:"Butorphanol"}
CAMP_LIGANDS = [l for l in FULL_PALETTE if l not in ("Naltrexone","C6Guano")]
def camp_model(x, top, bot, le): return bot + (top - bot) / (1 + 10.0**(le - x))
def parse_camp():
    ws = openpyxl.load_workbook(CAMP_XLSX, data_only=True)["Sheet1"]
    starts = sorted(CAMP_BLOCKS); data = {}
    for i, c in enumerate(starts):
        nxt = starts[i+1] if i+1 < len(starts) else 119
        pts = []
        for r in range(2, ws.max_row+1):
            xa = ws.cell(r, 1).value
            if xa in (None, ""): continue
            for cc in range(c, nxt):
                v = ws.cell(r, cc).value
                if v in (None, ""): continue
                fl = isinstance(v, str) and "*" in v
                try: fv = float(str(v).replace("*", "").strip())
                except ValueError: continue
                pts.append((float(xa), fv, fl))
        data[CAMP_BLOCKS[c]] = pts
    return data
def draw_camp(ax):
    data = parse_camp()
    fits = {}
    for lig in CAMP_LIGANDS:
        good = [(x, y) for (x, y, fl) in data[lig] if not fl and (lig, x) not in TP_EXCLUDE]
        x = np.array([p[0] for p in good]); y = np.array([p[1] for p in good])
        p, _ = curve_fit(camp_model, x, y, p0=[np.min(y), np.max(y), -8.0], maxfev=40000)
        fits[lig] = dict(top=p[0], bot=p[1], le=p[2])
    span = fits["DAMGO"]["bot"] - fits["DAMGO"]["top"]
    eff = lambda resp, bot: (bot - resp) / span * 100.0
    emax = {l: (fits[l]["bot"] - fits[l]["top"]) / span * 100 for l in CAMP_LIGANDS}
    order = ["DAMGO"] + [l for l in sorted(CAMP_LIGANDS, key=lambda l: -emax[l]) if l != "DAMGO"]
    hi, lo = [], []
    for lig in order:
        f = fits[lig]; col = color_for(lig); z = 10 if lig == "DAMGO" else 2
        good = [(x, y) for (x, y, fl) in data[lig] if not fl and (lig, x) not in TP_EXCLUDE]
        doses = sorted(set(x for x, _ in good))
        mx, my, se = [], [], []
        for d in doses:
            ys = np.array([y for x, y in good if x == d]); m = ys.mean()
            s = ys.std(ddof=1) / np.sqrt(len(ys)) if len(ys) > 1 else 0.0
            mx.append(d); my.append(eff(m, f["bot"])); se.append(s / span * 100)
        xlo = max(-14, min(doses) - 0.5); xhi = min(-4.5, max(doses) + 0.5)
        xx = np.linspace(xlo, xhi, 300)
        ax.plot(xx, eff(camp_model(xx, f["top"], f["bot"], f["le"]), f["bot"]), color=col, lw=1.0, zorder=z)
        ax.errorbar(mx, my, yerr=se, fmt="o", ms=2.2, mfc=col, mec="none",
                    ecolor=col, elinewidth=0.5, capsize=1, capthick=0.5, zorder=z + 1)
        hi += [a + b for a, b in zip(my, se)]; lo += [a - b for a, b in zip(my, se)]
    ax.set_xlim(-14, -4.5); ax.set_xticks(range(-14, -4, 1))
    ax.set_xticklabels([str(v) if v % 2 == 0 else "" for v in range(-14, -4, 1)])
    return hi, lo

# ── TruPath panel (RAW data, peak-anchored to match the Emax bar / Data 4) ────
# activity = 1 - value/own_-13_baseline ; point(dose) = per-dose mean / DAMGO_peak * 100.
# Each curve PLATEAUS at its peak Emax (= the bar value): tp_up(x)=Emax/(1+10^(le-x)),
# le from a per-ligand Hill=1 raw fit. This peak recipe reproduces Data 4 exactly, so
# the biased partials (PZM21 78.7, TRV130 77.2) top out at their reported efficacy and
# their supra-max points fall below the plateau (vs a span fit that under-read them).
TP_NAMES = ["DAMGO","PZM21","Methadone","Naloxone","Naltrexone","Carfentanil","Morphine",
            "Nalbuphine","Butorphanol","Buprenorphine","C6Guano","TRV130","MP","Fentanyl"]
TP_EXCLUDE = {("MP", -4.0)}   # highest-conc MP point is an outlier -> drop from fit + plot
def tp_up(x, emax, le): return emax / (1 + 10.0**(le - x))    # anchored plateau=emax, starts at 0
def parse_tp_raw():
    """{ligand: [replicate, ...]} where replicate = [(dose, value, flagged), ...]."""
    txt = open(TP_PZFX, encoding="utf-8", errors="replace").read()
    strip = lambda s: re.sub(r"<[^>]+>", "", s).strip()
    tabs = re.split(r"(?=<Table )", txt)
    rt = next(tb for tb in tabs if re.search(r"<Title>(.*?)</Title>", tb, re.S)
              and strip(re.search(r"<Title>(.*?)</Title>", tb, re.S).group(1)) == "20260117 µOR WT Gi1 TRUPATH")
    def subs(cb):
        return [[(strip(d) if d != "" else "") for d in re.findall(r"<d[^>]*/>|<d[^>]*>(.*?)</d>", s, re.S)]
                for s in re.findall(r"<Subcolumn>(.*?)</Subcolumn>", cb, re.S)]
    Xr = subs(re.search(r"<XColumn.*?</XColumn>", rt, re.S).group(0))[0]
    data = {}
    for yc in re.findall(r"<YColumn.*?</YColumn>", rt, re.S):
        nm = strip(re.search(r"<Title>(.*?)</Title>", yc, re.S).group(1))
        if nm not in FULL_PALETTE:
            continue
        reps = []
        for sub in subs(yc):
            rep = []
            for k, v in enumerate(sub):
                if k < len(Xr) and v != "" and Xr[k] != "":
                    fl = "*" in v
                    try: fv = float(v.replace("*", "").strip())
                    except ValueError: continue
                    x = float(Xr[k])
                    if (nm, x) in TP_EXCLUDE: continue
                    rep.append((x, fv, fl))
            reps.append(rep)
        data[nm] = reps
    return data
def tp_activity(reps):
    """Per-dose mean & SEM of activity (1 - value/own_-13_baseline), replicate-wise."""
    per = {}
    for rep in reps:
        base = next((y for (x, y, fl) in rep if x == -13.0 and not fl), None)
        if base in (None, 0):
            continue
        for (x, y, fl) in rep:
            if fl:
                continue
            per.setdefault(x, []).append(1.0 - y / base)
    doses = sorted(per)
    mean = {d: float(np.mean(per[d])) for d in doses}
    sem  = {d: (float(np.std(per[d], ddof=1) / np.sqrt(len(per[d]))) if len(per[d]) > 1 else 0.0) for d in doses}
    return mean, sem
def draw_tp(ax):
    data = parse_tp_raw()
    le, act = {}, {}
    for lig in TP_NAMES:
        pts = [(x, y) for rep in data[lig] for (x, y, fl) in rep if not fl]
        x = np.array([p[0] for p in pts]); y = np.array([p[1] for p in pts])
        p, _ = curve_fit(camp_model, x, y, p0=[np.min(y), np.max(y), -8.0], maxfev=40000)
        le[lig] = p[2]
        act[lig] = tp_activity(data[lig])
    dpk = max(act["DAMGO"][0].values())                     # DAMGO peak mean activity
    emax = {l: max(act[l][0].values()) / dpk * 100.0 for l in TP_NAMES}   # = Data 4 / bar
    order = ["DAMGO"] + [l for l in sorted(TP_NAMES, key=lambda l: -emax[l]) if l != "DAMGO"]
    hi, lo = [], []
    for lig in order:
        col = color_for(lig); z = 10 if lig == "DAMGO" else 2
        mean, sem = act[lig]
        doses = [d for d in sorted(mean) if d > -13]
        mx = doses
        my = [mean[d] / dpk * 100.0 for d in doses]
        se = [sem[d] / dpk * 100.0 for d in doses]
        xlo = max(-12.5, min(doses) - 0.5); xhi = min(-3.5, max(doses) + 0.5)
        xx = np.linspace(xlo, xhi, 300)
        ax.plot(xx, tp_up(xx, emax[lig], le[lig]), color=col, lw=1.0, zorder=z)
        ax.errorbar(mx, my, yerr=se, fmt="o", ms=2.2, mfc=col, mec="none",
                    ecolor=col, elinewidth=0.5, capsize=1, capthick=0.5, zorder=z + 1)
        hi += [a + b for a, b in zip(my, se)]; lo += [a - b for a, b in zip(my, se)]
    ax.set_xlim(-12.5, -3.5); ax.set_xticks(range(-12, -3, 1))
    ax.set_xticklabels([str(v) if v % 2 == 0 else "" for v in range(-12, -3, 1)])
    return hi, lo

def main():
    fig = plt.figure(); fig.set_size_inches(165 / 25.4, 52 / 25.4)
    axL = fig.add_axes([0.079, 0.20, 0.265, 0.62])
    axR = fig.add_axes([0.479, 0.20, 0.265, 0.62])   # wide gap = white space between plots
    hiL, loL = draw_camp(axL)
    hiR, loR = draw_tp(axR)

    ylim = (min(-10, min(loL + loR) - 3), max(hiL + hiR) + 4)
    for ax in (axL, axR):
        ax.axhline(0, color="grey", lw=0.4, ls=(0, (3, 3)), zorder=0)
        ax.set_ylim(*ylim); ax.set_yticks([0, 25, 50, 75, 100])
        ax.set_xlabel("Log[agonist] (M)"); ax.set_ylabel("Activity (% DAMGO)")
        for sp in ("top", "right"): ax.spines[sp].set_visible(False)
    axL.set_title("Pooled cAMP inhibition assay"); axR.set_title("TRUPATH GI1 BRET assay")

    handles = [Line2D([0], [0], marker="o", linestyle="none", markersize=2.6,
                      markerfacecolor=color_for(l), markeredgecolor="none", label=l)
               for l in LEGEND_ORDER]
    fig.legend(handles=handles, loc="center left", bbox_to_anchor=(0.755, 0.5),
               handlelength=1.0, handletextpad=0.4, labelspacing=0.32,
               borderpad=0.0, frameon=False)
    fig.savefig(OUT, transparent=True); print("Saved:", OUT)

if __name__ == "__main__":
    main()
