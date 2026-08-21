#!/usr/bin/env python3
"""
Compute Emax (% DAMGO) and LogEC50 for cAMP and TruPath and write them to
camp_trupath_scatter_values.csv for camp_vs_trupath_plots_black.R.

  cAMP    : fit transcriptional_drc.xlsx (3-param Hill=1, drop flagged),
            Emax = (Bottom-Top)/DAMGO_span * 100  (span-based; reproduces the
            published cAMP values exactly).
  TruPath : Emax is PEAK-based so it AGREES with the Emax bar and the DRC curves
            (= Prism 'Data 4'): per replicate / own -13 baseline -> activity =
            1 - norm; Emax = peak per-dose mean activity, % of DAMGO's peak.
            LogEC50 = per-ligand Hill=1 fit of the raw BRET. MP -4 outlier dropped.
"""
import re
import csv
import numpy as np
import openpyxl
from scipy.optimize import curve_fit

def model(x, top, bot, le): return bot + (top - bot) / (1 + 10.0**(le - x))
def fit(pts):
    x = np.array([p[0] for p in pts]); y = np.array([p[1] for p in pts])
    p, _ = curve_fit(model, x, y, p0=[np.min(y), np.max(y), -8.0], maxfev=40000)
    return dict(top=p[0], bot=p[1], le=p[2])

# ── cAMP (transcriptional_drc.xlsx) ──────────────────────────────────────────
CAMP_BLOCKS = {2:"DAMGO",11:"Fentanyl",20:"Methadone",29:"Buprenorphine",38:"Morphine",
               47:"Naloxone",56:"PZM21",65:"Uninduced cells",74:"TRV130",83:"MP",
               92:"Carfentanil",101:"Nalbuphine",110:"Butorphanol"}
def camp_fits():
    ws = openpyxl.load_workbook("transcriptional_drc.xlsx", data_only=True)["Sheet1"]
    starts = sorted(CAMP_BLOCKS); out = {}
    for i, c in enumerate(starts):
        nxt = starts[i+1] if i+1 < len(starts) else 119
        name = CAMP_BLOCKS[c]
        if name == "Uninduced cells": continue
        pts = []
        for r in range(2, ws.max_row+1):
            xa = ws.cell(r, 1).value
            if xa in (None, ""): continue
            for cc in range(c, nxt):
                v = ws.cell(r, cc).value
                if v in (None, "") or (isinstance(v, str) and "*" in v): continue
                try: pts.append((float(xa), float(str(v).replace("*", "").strip())))
                except ValueError: continue
        out[name] = fit(pts)
    return out

# ── TruPath (raw pzfx first tab) ─────────────────────────────────────────────
# PEAK-based Emax (matches bar + DRC curves / Data 4); EC50 from the Hill=1 raw fit.
TP_EXCLUDE = {("MP", -4.0)}
def tp_fits():
    txt = open("20260117_mor_wt_trupath_compilation.pzfx", encoding="utf-8", errors="replace").read()
    strip = lambda s: re.sub(r"<[^>]+>", "", s).strip()
    tabs = re.split(r"(?=<Table )", txt)
    rt = next(tb for tb in tabs if re.search(r"<Title>(.*?)</Title>", tb, re.S)
              and strip(re.search(r"<Title>(.*?)</Title>", tb, re.S).group(1)) == "20260117 µOR WT Gi1 TRUPATH")
    def subs(cb):
        return [[(strip(d) if d != "" else "") for d in re.findall(r"<d[^>]*/>|<d[^>]*>(.*?)</d>", s, re.S)]
                for s in re.findall(r"<Subcolumn>(.*?)</Subcolumn>", cb, re.S)]
    Xr = subs(re.search(r"<XColumn.*?</XColumn>", rt, re.S).group(0))[0]
    reps_by = {}
    for yc in re.findall(r"<YColumn.*?</YColumn>", rt, re.S):
        nm = strip(re.search(r"<Title>(.*?)</Title>", yc, re.S).group(1))
        reps = []
        for sub in subs(yc):
            rep = []
            for k, v in enumerate(sub):
                if k < len(Xr) and v != "" and Xr[k] != "":
                    fl = "*" in v
                    try: xx = float(Xr[k]); yy = float(v.replace("*", "").strip())
                    except ValueError: continue
                    if (nm, xx) in TP_EXCLUDE: continue
                    rep.append((xx, yy, fl))
            reps.append(rep)
        reps_by[nm] = reps
    def peak(reps):                       # peak of per-dose mean activity (1 - value/own -13)
        per = {}
        for rep in reps:
            base = next((y for (x, y, fl) in rep if x == -13.0 and not fl), None)
            if base in (None, 0): continue
            for (x, y, fl) in rep:
                if fl: continue
                per.setdefault(x, []).append(1.0 - y / base)
        return max((float(np.mean(v)) for v in per.values()), default=None)
    out = {}
    for nm, reps in reps_by.items():
        flat = [(x, y) for rep in reps for (x, y, fl) in rep if not fl]
        pk = peak(reps)
        if len(flat) < 4 or pk is None: continue
        try: out[nm] = dict(le=fit(flat)["le"], peak=pk)
        except Exception: pass
    return out

def emax(f, span): return (f["bot"] - f["top"]) / span * 100.0

cf = camp_fits(); tf = tp_fits()
camp_span = cf["DAMGO"]["bot"] - cf["DAMGO"]["top"]
tp_peak   = tf["DAMGO"]["peak"]
shared = [l for l in cf if l in tf]                       # 12 ligands in both assays
order = ["DAMGO"] + sorted([l for l in shared if l != "DAMGO"])

rows = []
for l in order:
    rows.append(dict(ligand=l,
                     camp_pct=round(emax(cf[l], camp_span), 2),
                     trupath_pct=round(tf[l]["peak"] / tp_peak * 100.0, 2),
                     camp_ec50=round(cf[l]["le"], 3),
                     trupath_ec50=round(tf[l]["le"], 3)))

with open("camp_trupath_scatter_values.csv", "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=["ligand","camp_pct","trupath_pct","camp_ec50","trupath_ec50"])
    w.writeheader(); w.writerows(rows)

print("%-14s %8s %8s %9s %9s" % ("ligand","cAMP%","TruP%","cAMP_EC50","TruP_EC50"))
for r in rows:
    print("%-14s %8.1f %8.1f %9.2f %9.2f" % (r["ligand"], r["camp_pct"], r["trupath_pct"],
                                             r["camp_ec50"], r["trupath_ec50"]))
print("\nSaved camp_trupath_scatter_values.csv (%d ligands)" % len(rows))
